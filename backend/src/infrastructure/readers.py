import csv
import re
from contextlib import contextmanager
from decimal import Decimal
from email import policy
from email.parser import BytesParser
from fnmatch import fnmatchcase
from io import BytesIO, StringIO
from pathlib import Path
from typing import Iterator

from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl.utils.exceptions import InvalidFileException
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from src.domain.models import (
    Account,
    BankEntry,
    CheckoutEntry,
    Contract,
    DataPackError,
    Dataset,
    PartnerTerms,
    RemittanceEntry,
    SourceRef,
)
from src.infrastructure.parsing import (
    blank,
    cnpj,
    day,
    decimal_value,
    identifier,
    integer,
    ofx_day,
    optional_text,
    percentage,
    text,
)
from src.infrastructure.source_scope import require_data_dir, require_source_file

_ROLES = {
    "checkout": "01_orbe_relatorio_comissoes_*.csv",
    "contracts": "02_tabela_taxas_contratadas*.xlsx",
    "email": "03_email_ponte_cambio_relatorio_*.eml",
    "accounts": "04_plano_de_contas_fenix_erp*.csv",
    "aurora": "05a_extrato_banco_aurora_*.csv",
    "meridiano": "05b_extrato_banco_meridiano_*.ofx",
    # Ajustes tardios Orbe (mesmo layout do 01): parte obrigatória do checkout da competência.
    "checkout_adjustments": "06_orbe_ajustes_tardios_*.csv",
}
_CHECKOUT_HEADER = (
    "id_transacao", "data_movimento", "cod_estabelecimento", "nome_estabelecimento",
    "tipo_lancamento", "modalidade", "qtd_parcelas", "bandeira", "valor_bruto",
    "taxa_aplicada_pct", "valor_taxa", "valor_liquido", "id_referencia", "data_venda_original",
)
_CONTRACT_HEADER = (
    "Lojista", "CNPJ", "EC Orbe", "Débito", "Crédito à vista", "Créd. 2-6x",
    "Créd. 7-12x", "Pix", "Antecipação", "Spread remessa", "Vigência", "Obs",
)
_RATE_KEYS = ("DEBITO", "CREDITO A VISTA", "CREDITO 2-6X", "CREDITO 7-12X", "PIX", "ANTECIPACAO")
_REMITTANCE_HEADER = (
    "Nº Operação", "Data Contratação", "Data Liquidação", "CNPJ Cliente", "Cliente",
    "Moeda", "Valor ME", "Taxa Cliente", "PTAX Ref.", "Valor R$", "Spread R$",
    "Spread %", "Spread Estimado R$ (D-1)", "IOF R$", "Status Compliance", "Observação",
)
_ACCOUNT_HEADER = ("COD_CONTA", "DESCRICAO", "TIPO", "NATUREZA", "ACEITA_LANCTO", "COD_SUPERIOR", "DT_CRIACAO")
_AURORA_HEADER = ("Data", "Lançamento", "Documento", "Valor (R$)", "Saldo (R$)")


def _select_files(data_dir: Path) -> dict[str, Path]:
    require_data_dir(data_dir)
    try:
        files = [path for path in data_dir.iterdir() if path.is_file()]
    except OSError:
        raise DataPackError("Diretório de dados inexistente ou inacessível.") from None
    selected = {}
    for role, pattern in _ROLES.items():
        matches = sorted(path for path in files if fnmatchcase(path.name.lower(), pattern))
        if len(matches) != 1:
            raise DataPackError(
                f"Papel {role}: esperado exatamente um arquivo {pattern}; encontrados {len(matches)}."
            )
        selected[role] = require_source_file(matches[0])
    return selected


def _bytes(path: Path) -> bytes:
    require_source_file(path)
    try:
        return path.read_bytes()
    except OSError:
        raise DataPackError("Não foi possível ler o arquivo.", SourceRef(path.name, "arquivo")) from None


def _csv(path: Path, encoding: str) -> list[tuple[int, list[str]]]:
    reader = None
    try:
        content = _bytes(path).decode(encoding)
        reader = csv.reader(StringIO(content, newline=""), delimiter=";", strict=True)
        rows = []
        previous_line = 0
        for row in reader:
            rows.append((previous_line + 1, row))
            previous_line = reader.line_num
        return rows
    except (UnicodeError, csv.Error):
        row = reader.line_num if reader is not None else None
        raise DataPackError("CSV inválido ou codificação incorreta.", SourceRef(path.name, f"L{row}" if row else "arquivo", row)) from None


def _shape(row: tuple | list, size: int, source: SourceRef) -> None:
    if len(row) != size:
        raise DataPackError(f"Esperadas {size} colunas; encontradas {len(row)}.", source)


def _header(row: tuple | list, expected: tuple[str, ...], source: SourceRef) -> None:
    if tuple(row) != expected:
        raise DataPackError("Cabeçalho não corresponde ao formato esperado.", source)


@contextmanager
def _workbook(payload: bytes, source: SourceRef) -> Iterator[Workbook]:
    with BytesIO(payload) as stream:
        try:
            workbook = load_workbook(stream, read_only=False, data_only=True, keep_links=False)
        except (BadZipFile, InvalidFileException, ParseError, EOFError, KeyError, ValueError, TypeError) as exc:
            raise DataPackError("Planilha XLSX inválida ou ilegível.", source) from exc
        try:
            yield workbook
        except (BadZipFile, InvalidFileException, ParseError, EOFError) as exc:
            raise DataPackError("Planilha XLSX inválida ou ilegível.", source) from exc
        finally:
            workbook.close()


def _sheet_rows(workbook: Workbook, sheet: str, source: SourceRef) -> list[tuple]:
    if sheet not in workbook.sheetnames:
        raise DataPackError(f"Aba obrigatória ausente: {sheet}.", source)
    return list(workbook[sheet].iter_rows(values_only=True))


def _checkout(path: Path) -> tuple[CheckoutEntry, ...]:
    rows = _csv(path, "utf-8-sig")
    if not rows:
        raise DataPackError("CSV sem cabeçalho.", SourceRef(path.name, "L1", 1))
    _header(rows[0][1], _CHECKOUT_HEADER, SourceRef(path.name, "L1", 1))
    entries = []
    for number, row in rows[1:]:
        if all(blank(value) for value in row):
            continue
        source = SourceRef(path.name, row[0].strip() or f"L{number}", number)
        _shape(row, len(_CHECKOUT_HEADER), source)
        transaction = text(row[0], source, "id_transacao")
        source = SourceRef(path.name, transaction, number)
        kind = text(row[4], source, "tipo_lancamento")
        modality = optional_text(row[5], source, "modalidade")
        installments = None if blank(row[6]) else integer(row[6], source, "qtd_parcelas")
        if modality is None and kind not in {"ANTECIPACAO", "REVERSAO ANTECIPACAO"}:
            raise DataPackError("modalidade: campo obrigatório para este lançamento.", source)
        entries.append(CheckoutEntry(
            source=source,
            occurred_on=day(row[1], source, "data_movimento"),
            establishment_code=identifier(row[2], source, "cod_estabelecimento"),
            merchant=text(row[3], source, "nome_estabelecimento"),
            kind=kind,
            modality=modality or "",
            installments=installments,
            gross=decimal_value(row[8], source, "valor_bruto", ptbr=True),
            applied_rate=percentage(row[9], source, "taxa_aplicada_pct", percentage_points=True),
            fee=decimal_value(row[10], source, "valor_taxa", ptbr=True),
            net=decimal_value(row[11], source, "valor_liquido", ptbr=True),
            original_ref=optional_text(row[12], source, "id_referencia"),
            original_on=None if blank(row[13]) else day(row[13], source, "data_venda_original"),
        ))
    return tuple(entries)


def _contracts(path: Path) -> tuple[Contract, ...]:
    sheet = "Taxas vigentes"
    base = SourceRef(path.name, "aba", sheet=sheet)
    with _workbook(_bytes(path), base) as workbook:
        rows = _sheet_rows(workbook, sheet, base)
        if len(rows) < 5:
            raise DataPackError("Planilha sem cabeçalho na linha 5.", base)
        _header(rows[4], _CONTRACT_HEADER, SourceRef(path.name, "A5:L5", 5, sheet))
        contracts = []
        for number, row in enumerate(rows[5:], 6):
            source = SourceRef(path.name, f"A{number}:L{number}", number, sheet)
            if all(blank(value) for value in row):
                continue
            _shape(row, len(_CONTRACT_HEADER), source)
            if row[0] == "* spread em branco = padrão" and all(blank(value) for value in row[1:]):
                continue
            rates = {
                key: percentage(row[index], source, _CONTRACT_HEADER[index])
                for index, key in enumerate(_RATE_KEYS, 3)
                if not blank(row[index])
            }
            remittance = row[9]
            default_remittance = blank(remittance) or isinstance(remittance, str) and remittance.strip().casefold() == "padrão"
            contracts.append(Contract(
                source=source,
                merchant=text(row[0], source, "Lojista"),
                cnpj=cnpj(row[1], source),
                establishment_code=None if blank(row[2]) else identifier(row[2], source, "EC Orbe"),
                rates=rates,
                remittance_rate=None if default_remittance else percentage(remittance, source, "Spread remessa"),
                validity=optional_text(row[10], source, "Vigência"),
                note=optional_text(row[11], source, "Obs"),
                comments="; ".join(cell.comment.text for cell in workbook[sheet][number] if cell.comment) or None,
            ))
    return tuple(contracts)


def _remittances(path: Path) -> tuple[RemittanceEntry, ...]:
    base = SourceRef(path.name, "MIME")
    message = BytesParser(policy=policy.default).parsebytes(_bytes(path))
    if any(part.defects for part in message.walk()):
        raise DataPackError("E-mail MIME malformado.", base)
    attachments = [
        part for part in message.walk()
        if part.get_filename()
        and fnmatchcase(part.get_filename().lower(), "pontecambio_relatorio_tarvo_*.xlsx")
    ]
    if len(attachments) != 1:
        raise DataPackError(f"Esperado exatamente um anexo PonteCambio XLSX; encontrados {len(attachments)}.", base)
    attachment = attachments[0]
    filename = attachment.get_filename().replace("\\", "/").rsplit("/", 1)[-1]
    origin = f"{path.name}!{filename}"
    base = SourceRef(origin, "anexo", sheet="Operações")
    payload = attachment.get_payload(decode=True)
    if not payload or attachment.defects:
        raise DataPackError("Anexo XLSX ausente ou codificação MIME inválida.", base)
    with _workbook(payload, base) as workbook:
        rows = _sheet_rows(workbook, "Operações", base)
        if len(rows) < 5:
            raise DataPackError("Operações sem cabeçalho na linha 5.", base)
        _header(rows[4], _REMITTANCE_HEADER, SourceRef(origin, "A5:P5", 5, "Operações"))
        entries = []
        for number, row in enumerate(rows[5:], 6):
            source = SourceRef(origin, f"A{number}:P{number}", number, "Operações")
            if all(blank(value) for value in row):
                continue
            _shape(row, len(_REMITTANCE_HEADER), source)
            if row[0] in {"Total de operações:", "Spread total R$:"}:
                if not all(blank(value) for value in row[2:]):
                    raise DataPackError("Rodapé de operações malformado.", source)
                if row[0] == "Total de operações:":
                    integer(row[1], source, "Total de operações")
                else:
                    decimal_value(row[1], source, "Spread total R$")
                continue
            operation = text(row[0], source, "Nº Operação")
            if not operation.startswith("PC-"):
                raise DataPackError("Nº Operação: identificador PC- obrigatório.", source)
            source = SourceRef(origin, operation, number, "Operações")
            entries.append(RemittanceEntry(
                source=source,
                contracted_on=day(row[1], source, "Data Contratação"),
                settled_on=None if blank(row[2]) else day(row[2], source, "Data Liquidação"),
                cnpj=cnpj(row[3], source),
                merchant=text(row[4], source, "Cliente"),
                currency=text(row[5], source, "Moeda"),
                foreign_amount=decimal_value(row[6], source, "Valor ME"),
                customer_rate=decimal_value(row[7], source, "Taxa Cliente"),
                ptax_rate=decimal_value(row[8], source, "PTAX Ref."),
                brl_amount=decimal_value(row[9], source, "Valor R$"),
                spread=decimal_value(row[10], source, "Spread R$"),
                spread_rate=percentage(row[11], source, "Spread %"),
                estimated_spread=decimal_value(row[12], source, "Spread Estimado R$ (D-1)"),
                iof=decimal_value(row[13], source, "IOF R$"),
                status=text(row[14], source, "Status Compliance"),
                note=optional_text(row[15], source, "Observação"),
            ))
    return tuple(entries)


def _accounts(path: Path) -> tuple[Account, ...]:
    rows = _csv(path, "cp1252")
    if not rows:
        raise DataPackError("Plano de contas sem cabeçalho.", SourceRef(path.name, "L1", 1))
    _header(rows[0][1], _ACCOUNT_HEADER, SourceRef(path.name, "L1", 1))
    accounts = []
    for number, row in rows[1:]:
        if all(blank(value) for value in row):
            continue
        source = SourceRef(path.name, row[0].strip() or f"L{number}", number)
        _shape(row, len(_ACCOUNT_HEADER), source)
        code = text(row[0], source, "COD_CONTA")
        source = SourceRef(path.name, code, number)
        posting = text(row[4], source, "ACEITA_LANCTO")
        if posting not in {"S", "N"}:
            raise DataPackError("ACEITA_LANCTO: esperado S ou N.", source)
        day(row[6], source, "DT_CRIACAO")
        accounts.append(Account(
            source=source,
            code=code,
            description=text(row[1], source, "DESCRICAO"),
            kind=text(row[2], source, "TIPO"),
            nature=text(row[3], source, "NATUREZA"),
            posting_allowed=posting == "S",
        ))
    return tuple(accounts)


def _aurora(path: Path) -> tuple[tuple[BankEntry, ...], dict[str, Decimal], dict[str, Decimal]]:
    rows = _csv(path, "utf-8-sig")
    if len(rows) < 8:
        raise DataPackError("Extrato Aurora sem cabeçalho ou saldo anterior.", SourceRef(path.name, "L7", 7))
    _header(rows[6][1], _AURORA_HEADER, SourceRef(path.name, "L7", 7))
    opening_number, opening_row = rows[7]
    source = SourceRef(path.name, f"L{opening_number}", opening_number)
    _shape(opening_row, len(_AURORA_HEADER), source)
    if opening_row[1] != "SALDO ANTERIOR" or any(not blank(value) for value in opening_row[2:4]):
        raise DataPackError("Saldo anterior Aurora inválido.", source)
    day(opening_row[0], source, "Data saldo anterior")
    opening = decimal_value(opening_row[4], source, "Saldo anterior", ptbr=True)
    closing = None
    entries = []
    for number, row in rows[8:]:
        source = SourceRef(path.name, f"L{number}", number)
        if all(blank(value) for value in row):
            continue
        _shape(row, len(_AURORA_HEADER), source)
        if row[1] == "SALDO FINAL":
            if closing is not None or any(not blank(row[index]) for index in (0, 2, 3)):
                raise DataPackError("Saldo final Aurora inválido ou duplicado.", source)
            closing = decimal_value(row[4], source, "Saldo final", ptbr=True)
            continue
        if row[1] == "Os lançamentos estão sujeitos a confirmação.":
            if closing is None or any(not blank(row[index]) for index in (0, 2, 3, 4)):
                raise DataPackError("Rodapé Aurora malformado.", source)
            continue
        if closing is not None:
            raise DataPackError("Lançamento encontrado depois do saldo final Aurora.", source)
        entries.append(BankEntry(
            source=source,
            bank="AURORA",
            occurred_on=day(row[0], source, "Data"),
            description=text(row[1], source, "Lançamento"),
            amount=decimal_value(row[3], source, "Valor (R$)", ptbr=True),
            balance=decimal_value(row[4], source, "Saldo (R$)", ptbr=True),
        ))
    if closing is None:
        raise DataPackError("Saldo final Aurora ausente.", SourceRef(path.name, "SALDO FINAL"))
    return tuple(entries), {"AURORA": opening}, {"AURORA": closing}


def _ofx_value(block: str, tag: str, source: SourceRef) -> str:
    values = re.findall(rf"<{tag}>\s*([^<]*)", block, flags=re.IGNORECASE)
    if len(values) != 1 or not values[0].strip():
        raise DataPackError(f"OFX: campo {tag} ausente, vazio ou duplicado.", source)
    return values[0].strip()


def _ofx_block(content: str, tag: str, source: SourceRef) -> str:
    blocks = re.findall(rf"<{tag}>(.*?)</{tag}>", content, flags=re.IGNORECASE | re.DOTALL)
    if len(blocks) != 1:
        raise DataPackError(f"OFX: bloco {tag} ausente ou ambíguo.", source)
    return blocks[0]


def _meridiano(path: Path) -> tuple[tuple[BankEntry, ...], dict[str, Decimal]]:
    base = SourceRef(path.name, "OFX")
    try:
        content = _bytes(path).decode("cp1252")
    except UnicodeError:
        raise DataPackError("OFX: codificação CP1252 inválida.", base) from None
    root = _ofx_block(content, "OFX", base)
    transactions = _ofx_block(root, "BANKTRANLIST", base)
    blocks = list(re.finditer(r"<STMTTRN>(.*?)</STMTTRN>", transactions, flags=re.IGNORECASE | re.DOTALL))
    opened = len(re.findall(r"<STMTTRN>", transactions, flags=re.IGNORECASE))
    closed = len(re.findall(r"</STMTTRN>", transactions, flags=re.IGNORECASE))
    if len(blocks) != opened or opened != closed:
        raise DataPackError("OFX: lançamento STMTTRN incompleto.", base)
    ofx_day(_ofx_value(transactions, "DTSTART", base), base, "DTSTART")
    ofx_day(_ofx_value(transactions, "DTEND", base), base, "DTEND")
    entries = []
    for match in blocks:
        block = match[1]
        fitid = _ofx_value(block, "FITID", base)
        source = SourceRef(path.name, fitid)
        entries.append(BankEntry(
            source=source,
            bank="MERIDIANO",
            occurred_on=ofx_day(_ofx_value(block, "DTPOSTED", source), source, "DTPOSTED"),
            description=_ofx_value(block, "MEMO", source),
            amount=decimal_value(_ofx_value(block, "TRNAMT", source), source, "TRNAMT"),
        ))
    ledger = _ofx_block(root, "LEDGERBAL", base)
    source = SourceRef(path.name, "LEDGERBAL")
    closing = decimal_value(_ofx_value(ledger, "BALAMT", source), source, "BALAMT")
    ofx_day(_ofx_value(ledger, "DTASOF", source), source, "DTASOF")
    return tuple(entries), {"MERIDIANO": closing}


def _partner_terms(path: Path) -> PartnerTerms:
    message = BytesParser(policy=policy.default).parsebytes(_bytes(path))
    body = message.get_body(preferencelist=("plain",))
    content = body.get_content() if body else ""
    participation = re.findall(r"participa[çc][ãa]o de\s+(\d+(?:[.,]\d+)?)%", content, re.IGNORECASE)
    signature = re.findall(r"([^\n]+?)\s*-\s*CNPJ\s+([\d./-]+)", content)
    name, taxpayer = signature[0] if len(signature) == 1 else (None, None)
    if name:
        name = name.rsplit(" - ", 1)[-1].strip()
    source = SourceRef(path.name, "MIME:text/plain")
    return PartnerTerms(source,
        Decimal(participation[0].replace(",", ".")) / 100 if len(participation) == 1 else None,
        cnpj(taxpayer, source) if taxpayer else None, name)


def load_dataset(data_dir: Path) -> Dataset:
    files = _select_files(data_dir)
    checkout = _checkout(files["checkout"]) + _checkout(files["checkout_adjustments"])
    contracts = _contracts(files["contracts"])
    remittances = _remittances(files["email"])
    accounts = _accounts(files["accounts"])
    aurora, opening, closing = _aurora(files["aurora"])
    meridiano, meridiano_closing = _meridiano(files["meridiano"])
    return Dataset(
        checkout=checkout,
        contracts=contracts,
        remittances=remittances,
        banks=aurora + meridiano,
        accounts=accounts,
        bank_opening_balances=opening,
        bank_closing_balances=closing | meridiano_closing,
        partner_terms=_partner_terms(files["email"]),
    )
