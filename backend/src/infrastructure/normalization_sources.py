import hashlib
import json
import re
import unicodedata
from dataclasses import asdict
from datetime import date, datetime
from decimal import Decimal
from email import policy
from email.parser import BytesParser
from pathlib import Path

from openpyxl import load_workbook
from rapidfuzz.fuzz import WRatio

from src.domain.models import Contract, Dataset
from src.domain.policies import ESTABLISHMENT_BINDINGS


def serialize(value: object) -> str:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    raise TypeError(f"Tipo não serializável: {type(value).__name__}")


def json_text(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, default=serialize)


def digest(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def normalize_name(value: str) -> str:
    plain = unicodedata.normalize("NFKD", value.casefold())
    return " ".join(re.sub(r"[^a-z0-9]+", " ", "".join(c for c in plain if not unicodedata.combining(c))).split())


def similarity(left: str, right: str) -> float:
    return WRatio(normalize_name(left), normalize_name(right))


def contract_id(contract: Contract) -> str:
    return f"{contract.source.file}!{contract.source.sheet}!{contract.source.ref}"


def contract_rows(data_dir: Path, dataset: Dataset) -> list[dict]:
    files = {contract.source.file for contract in dataset.contracts}
    if len(files) != 1:
        raise ValueError("Esperado um arquivo de contratos para normalização.")
    path = data_dir / next(iter(files))
    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    records = []
    try:
        for sheet in workbook:
            historical = sheet.title != "Taxas vigentes"
            first_row = 2 if historical else 6
            for cells in sheet.iter_rows(min_row=first_row):
                if not cells[0].value or not cells[1].value:
                    continue
                row = cells[0].row
                source_id = f"{path.name}!{sheet.title}!A{row}:{sheet.cell(row, sheet.max_column).column_letter}{row}"
                records.append({
                    "source_id": source_id, "historical": historical,
                    "merchant": str(cells[0].value),
                    "cnpj": re.sub(r"\D", "", str(cells[1].value)).zfill(14),
                    "establishment_code": None if historical or cells[2].value is None else str(cells[2].value),
                    "validity": None if historical else cells[10].value,
                    "note": None if historical else cells[11].value,
                    "cells": {c.coordinate: {"value": str(c.value) if isinstance(c.value, float) else c.value,
                        "format": c.number_format, "type": c.data_type,
                        "comment": c.comment.text if c.comment else None} for c in cells},
                })
    finally:
        workbook.close()
    current = {contract_id(c) for c in dataset.contracts}
    if {r["source_id"] for r in records if not r["historical"]} != current:
        raise ValueError("As linhas extraídas não correspondem aos contratos do leitor financeiro.")
    return records


def related_rows(focal: dict, rows: list[dict], threshold: float, limit: int) -> list[dict]:
    ranked = []
    for row in rows:
        exact = row["cnpj"] == focal["cnpj"] or bool(
            focal["establishment_code"] and row["establishment_code"] == focal["establishment_code"]
        )
        score = similarity(focal["merchant"], row["merchant"])
        if exact or score >= threshold:
            ranked.append((not exact, -score, row["source_id"], row))
    return [item[-1] for item in sorted(ranked)[:limit]]


def orbe_candidates(focal: dict, dataset: Dataset, threshold: float, sample_limit: int) -> list[dict]:
    groups: dict[str, list] = {}
    for entry in dataset.checkout:
        groups.setdefault(entry.establishment_code, []).append(entry)
    selected = []
    for ec, entries in sorted(groups.items()):
        score = max(similarity(focal["merchant"], entry.merchant) for entry in entries)
        confirmed = ESTABLISHMENT_BINDINGS.get(ec) == focal["cnpj"]
        if ec != focal["establishment_code"] and not confirmed and score < threshold:
            continue
        ordered = sorted(entries, key=lambda e: (e.occurred_on, e.source.ref))
        selected.append({
            "establishment_code": ec, "names": sorted({e.merchant for e in entries}),
            "name_score": round(score, 2), "confirmed_cnpj": ESTABLISHMENT_BINDINGS.get(ec),
            "confirmation_source": "policies.ESTABLISHMENT_BINDINGS; financeiro thread 4" if confirmed else None,
            "total_rows": len(entries), "samples": [asdict(e) for e in ordered[:sample_limit]],
        })
    return selected


def document_chunks(docs_dir: Path) -> list[dict]:
    chunks = []
    if not docs_dir.exists():
        return chunks
    for path in sorted(docs_dir.glob("*.md")):
        lines = path.read_text(encoding="utf-8-sig").splitlines()
        for start in range(0, len(lines), 12):
            chunks.append({"source_id": f"{path.name}:L{start + 1}-L{min(start + 12, len(lines))}",
                "file": path.name, "text": "\n".join(lines[start:start + 12])})
    return chunks

def email_chunks(data_dir: Path) -> list[dict]:
    chunks = []
    if not data_dir.exists():
        return chunks
    for path in sorted(data_dir.glob("*.eml")):
        try:
            message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        except Exception:
            continue
        headers = []
        for header in ("From", "To", "Cc", "Date", "Subject"):
            if message[header]:
                headers.append(f"{header}: {message[header]}")
        if headers:
            chunks.append({
                "source_id": f"{path.name}:headers",
                "file": path.name,
                "text": "\n".join(headers),
            })
        body_parts = []
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                payload = part.get_payload(decode=True)
                if payload:
                    charset = part.get_content_charset() or "utf-8"
                    body_parts.append(payload.decode(charset, errors="replace"))
            elif part.get_content_type() == "text/html" and not body_parts:
                payload = part.get_payload(decode=True)
                if payload:
                    charset = part.get_content_charset() or "utf-8"
                    raw = payload.decode(charset, errors="replace")
                    body_parts.append(re.sub(r"<[^>]+>", " ", raw))
        for part_idx, body in enumerate(body_parts, 1):
            paragraphs = [p.strip() for p in body.split("\n\n") if p.strip()]
            for p_idx, paragraph in enumerate(paragraphs, 1):
                chunks.append({
                    "source_id": f"{path.name}:part{part_idx}:P{p_idx}",
                    "file": path.name,
                    "text": paragraph,
                })
    return chunks


def retrieve_documents(chunks: list[dict], focal: dict, ec: str, limit: int = 24) -> list[dict]:
    terms = {focal["cnpj"], ec, normalize_name(focal["merchant"])} - {""}
    ranked = []
    for chunk in chunks:
        plain = normalize_name(chunk["text"])
        digits = re.sub(r"\D", "", chunk["text"])
        score = sum(5 for term in terms if term in plain or term in digits)
        score += 2 * (chunk["file"] == "THREADS_FINANCEIRO_TARVO.md")
        score += 1 * (chunk["file"] == "HIPOTESES.md")
        score += 3 * (chunk["file"].endswith(".eml") and any(term in plain or term in digits for term in terms))
        if score:
            ranked.append((-score, chunk["source_id"], chunk))
    return [item[-1] for item in sorted(ranked)[:limit]]


def input_manifest(data_dir: Path, docs_dir: Path) -> dict[str, str]:
    files = sorted(p for p in data_dir.iterdir() if p.is_file() and not p.name.startswith("."))
    documents = sorted(docs_dir.glob("*.md")) if docs_dir.exists() else []
    return {f"dados/{p.name}": digest(p.read_bytes()) for p in files} | {
        f"docs/{p.name}": digest(p.read_bytes()) for p in documents
    }
